"""
Export company profiles to Excel.

Layout — one sheet per company:
  - Columns = questions
  - Row 1   = Final Answer
  - Row 2,3 = Excerpt 1, Source 1 (clickable link)
  - Row 4,5 = Excerpt 2, Source 2
  - ...     = up to the max number of sources for any question in that company

Cells with no source for a given row are left blank.
"""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_LINK_FONT = Font(color="0563C1", underline="single")
_BOLD = Font(bold=True)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")


def export_profiles_to_excel(profiles: list[dict], output_path: Path) -> None:
    """Write one Excel workbook with one sheet per company."""
    if not profiles:
        return

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for profile in profiles:
            _write_company_sheet(writer, profile)


def _write_company_sheet(writer: pd.ExcelWriter, profile: dict) -> None:
    company = profile["company"]
    answers = profile.get("answers", [])

    max_sources = max((len(a.get("sources", [])) for a in answers), default=0)

    row_labels = ["Final Answer"]
    for i in range(1, max_sources + 1):
        row_labels.append(f"Excerpt {i}")
        row_labels.append(f"Source {i}")

    data: dict[str, list[str]] = {}
    for a in answers:
        sources = a.get("sources", [])
        col = [a.get("answer", "")]
        for i in range(max_sources):
            if i < len(sources):
                col.append(sources[i].get("excerpt", ""))
                col.append(sources[i].get("url", ""))
            else:
                col.append("")
                col.append("")
        data[a["question"]] = col

    df = pd.DataFrame(data, index=row_labels)
    sheet_name = _safe_sheet_name(company)
    df.to_excel(writer, sheet_name=sheet_name)

    _format_sheet(writer.sheets[sheet_name], df)


def _safe_sheet_name(name: str) -> str:
    for ch in ['\\', '/', '*', '[', ']', ':', '?']:
        name = name.replace(ch, '')
    return (name or "Sheet")[:31]


def _format_sheet(ws, df: pd.DataFrame) -> None:
    n_cols = len(df.columns)
    n_rows = len(df.index)

    # Header row (questions)
    for col_idx in range(1, n_cols + 2):
        cell = ws.cell(row=1, column=col_idx)
        cell.alignment = _WRAP_TOP
        cell.font = _BOLD

    # Row labels (column A) + data cells
    for row_idx in range(2, n_rows + 2):
        label_cell = ws.cell(row=row_idx, column=1)
        label_cell.font = _BOLD
        label_cell.alignment = _WRAP_TOP
        label = str(label_cell.value or "")
        is_source_row = label.startswith("Source")

        for col_idx in range(2, n_cols + 2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = _WRAP_TOP

            if is_source_row and cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = _LINK_FONT

        ws.row_dimensions[row_idx].height = 28 if is_source_row else 110

    # Column widths
    ws.column_dimensions["A"].width = 16
    for col_idx in range(2, n_cols + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 60

    ws.row_dimensions[1].height = 45

    # Freeze header row + label column for easier scrolling
    ws.freeze_panes = "B2"
